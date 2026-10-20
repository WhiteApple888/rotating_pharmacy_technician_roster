"""Excel workbook input for the fixed roster layout and RD history sheet."""

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .individual_staff_constraints import Roster, TOTAL_NUMBER_OF_STAFF, StaffSchedule
from .parsing import day_columns, normalise

CURRENT_ROSTER_START_ROW = 6

LAST_ROSTER_RD_SHEET = "lastRosterRD"
LAST_ROSTER_RD_COLUMN = 2  # Excel column B
LAST_ROSTER_RD_START_ROWS = 2

@dataclass(frozen=True)
class LastRosterRestDays:
    """The 16 prior-RD values, in the same order as ``STAFF_ROWS``."""

    dates: tuple[date, ...]

WorkbookSource = Path | BinaryIO | BytesIO


def load_roster(
    source: WorkbookSource,
    sheet_name: str | None,
    num_staff: int = TOTAL_NUMBER_OF_STAFF,
    start_row_current_roster: int = CURRENT_ROSTER_START_ROW,
) -> Roster:
    """Read a 28-day roster grid from a path or an in-memory workbook.

    The roster layout starts at Excel row 6. ``num_staff`` is configurable for
    API consumers, while the CLI retains its established 16-staff default.
    """
    if num_staff <= 0:
        raise ValueError("num_staff must be a positive integer.")

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        columns = day_columns()
        staff = tuple(
            StaffSchedule(row, tuple(normalise(sheet.cell(row=row, column=column).value) for column in columns))
            for row in range(start_row_current_roster, start_row_current_roster + num_staff)
        )
        return Roster(worksheet=sheet.title, staff=staff)
    except KeyError as exc:
        raise ValueError(f"Worksheet '{sheet_name}' does not exist.") from exc
    finally:
        workbook.close()


def load_last_roster_rest_days(
    source: WorkbookSource,
    start_date: date,
    num_staff: int = TOTAL_NUMBER_OF_STAFF,
    start_row_last_roster_rd: int = LAST_ROSTER_RD_START_ROWS,

) -> LastRosterRestDays:
    """Read and validate per-person history from ``lastRosterRD`` column B.

    ``B2`` maps to roster row 6 and subsequent cells map to subsequent roster
    rows.  The sheet must contain one Excel date per submitted staff member,
    and every date must be before Day 1.
    """
    if num_staff <= 0:
        raise ValueError("num_staff must be a positive integer.")

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        try:
            sheet = workbook[LAST_ROSTER_RD_SHEET]
        except KeyError as exc:
            raise ValueError(
                f"Worksheet '{LAST_ROSTER_RD_SHEET}' is required for the baseline RD check."
            ) from exc

        dates: list[date] = []
        for row in range(start_row_last_roster_rd, start_row_last_roster_rd + num_staff):
            cell = sheet.cell(row=row, column=LAST_ROSTER_RD_COLUMN)
            value = cell.value
            if isinstance(value, (int, float)):
                value = from_excel(value, workbook.epoch)

            if isinstance(value, datetime):
                value = value.date()

            if not isinstance(value, date):
                raise ValueError(
                    f"{LAST_ROSTER_RD_SHEET}!{cell.coordinate} must contain an Excel date; "
                    f"received {value!r}."
                )

            if value >= start_date:
                raise ValueError(
                    f"{LAST_ROSTER_RD_SHEET}!{cell.coordinate} must be before roster Day 1 "
                    f"({start_date.isoformat()}); received {value.isoformat()}."
                )

            dates.append(value)

        return LastRosterRestDays(tuple(dates))
    finally:
        workbook.close()
