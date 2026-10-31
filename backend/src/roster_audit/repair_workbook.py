"""Read and write the Excel representation of a roster repair."""

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .individual_policy import SHIFT_HOURS
from .individual_staff_constraints import Roster
from .optimization_policy import STAFF_COUNT, _change_cost, _marker, valid_takeover_priority
from .parsing import day_columns
from .workbook import CURRENT_ROSTER_START_ROW, load_last_roster_rest_days, load_roster

PRIORITY_COLUMNS = (3, 14, 25, 36)
WEEKLY_HOURS_COLUMNS = (11, 22, 33, 44)


def load_repair_input(workbook_bytes: bytes, start_date: date, worksheet_name: str | None,
                      num_staff: int = STAFF_COUNT, fixed_staff: frozenset[int] = frozenset()):
    """Read the roster, previous rest days, and weekly takeover priorities."""
    if worksheet_name is None:
        workbook = load_workbook(BytesIO(workbook_bytes), read_only=True)
        try:
            worksheet_name = next(
                (name for name in workbook.sheetnames if name not in {"lastRosterRD", "Change Log"}), None,
            )
        finally:
            workbook.close()
        if worksheet_name is None:
            raise ValueError("Workbook has no roster worksheet.")
    roster = load_roster(BytesIO(workbook_bytes), worksheet_name, num_staff=num_staff)
    if roster.worksheet == "lastRosterRD":
        raise ValueError("Select the roster worksheet, not lastRosterRD' worksheet.")
    prior = load_last_roster_rest_days(
        BytesIO(workbook_bytes), start_date, num_staff=num_staff,
    ).dates
    _validate_fixed_formulas(workbook_bytes, roster.worksheet, roster, fixed_staff)
    return roster, prior, _load_priorities(workbook_bytes, roster.worksheet, num_staff, fixed_staff)


def load_exported_roster(workbook_bytes: bytes, sheet_name: str, num_staff: int = STAFF_COUNT) -> Roster:
    """Read the serialized export afresh for independent verification."""
    return load_roster(BytesIO(workbook_bytes), sheet_name, num_staff=num_staff)


def _load_priorities(workbook_bytes: bytes, sheet_name: str, num_staff: int = STAFF_COUNT,
                     fixed_staff: frozenset[int] = frozenset()) -> tuple[dict[int, int], ...]:
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        weeks = []
        for week, column in enumerate(PRIORITY_COLUMNS, start=1):
            ranks: dict[int, int] = {}
            used: set[int] = set()
            for staff in range(1, num_staff + 1):
                value = sheet.cell(staff + CURRENT_ROSTER_START_ROW - 1, column).value
                if value is None or value == "":
                    continue
                if not valid_takeover_priority(value, staff, used, frozenset()):
                    raise ValueError(f"Week {week} has an invalid takeover priority at row {staff + CURRENT_ROSTER_START_ROW - 1}.")
                ranks[staff] = value
                used.add(value)
            weeks.append(ranks)
        return tuple(weeks)
    finally:
        workbook.close()


def write_repaired_roster(workbook_bytes: bytes, sheet_name: str, original: Roster,
                  roster: Roster, start_date: date, penalized_changes: int,
                  cell_edits: int, plus_changes: int, fixed_staff: frozenset[int] = frozenset()) -> bytes:
    workbook = load_workbook(BytesIO(workbook_bytes))
    try:
        sheet = workbook[sheet_name]
        columns = day_columns()
        for staff, schedule in enumerate(roster.staff, start=1):
            if staff not in fixed_staff:
                for column, code in zip(columns, schedule.shifts):
                    sheet.cell(schedule.excel_row, column).value = code
            for week, column in enumerate(WEEKLY_HOURS_COLUMNS):
                hours_cell = sheet.cell(schedule.excel_row, column)
                # Preserve formulas and intentionally blank cells. The source
                # workbook uses numeric snapshots for the first 15 staff.
                if isinstance(hours_cell.value, (int, float)) and not isinstance(
                    hours_cell.value, bool
                ):
                    hours_cell.value = sum(
                        SHIFT_HOURS[code]
                        for code in schedule.shifts[week * 7:(week + 1) * 7]
                    )
        if "Change Log" in workbook:
            del workbook["Change Log"]
        log = workbook.create_sheet("Change Log")
        for label, value in (("Penalized change count", penalized_changes),
                             ("Cell edit count", cell_edits),
                             ("Plus change count", plus_changes)):
            log.append([label, value])
        log.append([])
        log.append(["Excel row", "Column", "Date", "Staff name", "Before", "After",
                    "Base shift changed", "Plus marker changed", "Penalized cost changed"])
        for before, after in zip(original.staff, roster.staff):
            for day, (old, new) in enumerate(zip(before.shifts, after.shifts)):
                if old == new:
                    continue
                log.append([before.excel_row, get_column_letter(columns[day]),
                            start_date + timedelta(days=day),
                            sheet.cell(before.excel_row, 1).value, old, new,
                            old.rstrip("+") != new.rstrip("+"),
                            _marker(old) != _marker(new), bool(_change_cost(old, new))])
                log.cell(log.max_row, 3).number_format = "yyyy-mm-dd"
        for cell in log[5]:
            cell.font = Font(bold=True)
        log.freeze_panes = "A6"
        log.auto_filter.ref = f"A5:I{max(5, log.max_row)}"
        for column, width in enumerate((12, 10, 14, 22, 14, 14, 20, 21, 24), 1):
            log.column_dimensions[get_column_letter(column)].width = width
        output = BytesIO()
        workbook.save(output)
        return _restore_fixed_formula_caches(workbook_bytes, output.getvalue(), sheet_name, original, fixed_staff)
    finally:
        workbook.close()


def _validate_fixed_formulas(workbook_bytes: bytes, sheet_name: str, roster: Roster,
                             fixed_staff: frozenset[int]) -> None:
    """Require a readable calculated shift for each fixed formula cell."""
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name]
        for staff in fixed_staff:
            schedule = roster.staff[staff - 1]
            for column, code in zip(day_columns(), schedule.shifts):
                cell = sheet.cell(schedule.excel_row, column)
                if cell.data_type == "f" and code not in SHIFT_HOURS:
                    raise ValueError(
                        f"Fixed staff {staff} cell {cell.coordinate} has no readable calculated shift value; "
                        "recalculate and save the workbook before optimizing."
                    )
    finally:
        workbook.close()


def _restore_fixed_formula_caches(source: bytes, output: bytes, sheet_name: str,
                                   original: Roster, fixed_staff: frozenset[int]) -> bytes:
    """Preserve cached formula results removed by openpyxl when saving."""
    import re
    from html import escape
    from posixpath import normpath, join
    from xml.etree import ElementTree as ET
    from zipfile import ZipFile

    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "p": "http://schemas.openxmlformats.org/package/2006/relationships"}

    def worksheet_path(archive):
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationship_id = next(s.get(f"{{{ns['r']}}}id") for s in workbook.findall("m:sheets/m:sheet", ns)
                               if s.get("name") == sheet_name)
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(r.get("Target") for r in relations if r.get("Id") == relationship_id)
        return target.lstrip("/") if target.startswith("/") else normpath(join("xl", target))

    if not fixed_staff:
        return output
    coordinates = {f"{get_column_letter(column)}{original.staff[staff - 1].excel_row}"
                   for staff in fixed_staff for column in day_columns()}
    with ZipFile(BytesIO(source)) as source_zip, ZipFile(BytesIO(output)) as output_zip:
        source_path, output_path = worksheet_path(source_zip), worksheet_path(output_zip)
        source_root = ET.fromstring(source_zip.read(source_path))
        output_xml = output_zip.read(output_path).decode("utf-8")
        source_cells = {cell.get("r"): cell for cell in source_root.findall(".//m:sheetData/m:row/m:c", ns)
                        if cell.get("r") in coordinates and cell.find("m:f", ns) is not None}
        if not source_cells:
            return output
        for coordinate, source_cell in source_cells.items():
            cached = source_cell.find("m:v", ns)
            value = "" if cached is None or cached.text is None else escape(cached.text)
            cell_type = source_cell.get("t")
            pattern = rf'(<c\b[^>]*\br="{re.escape(coordinate)}"[^>]*>)(.*?)(</c>)'
            def restore(match):
                opening, body, closing = match.groups()
                opening = re.sub(r'\s+t="[^"]*"', '', opening)
                if cell_type:
                    opening = opening[:-1] + f' t="{escape(cell_type, quote=True)}">'
                body = re.sub(r'<v(?:\s[^>]*)?\s*/>|<v(?:\s[^>]*)?>.*?</v>', '', body)
                return opening + body + f"<v>{value}</v>" + closing
            output_xml, count = re.subn(pattern, restore, output_xml, count=1, flags=re.DOTALL)
            if count != 1:
                raise ValueError(f"Could not preserve fixed formula cell {coordinate}.")
        result = BytesIO()
        with ZipFile(result, "w") as result_zip:
            for info in output_zip.infolist():
                content = (output_xml.encode("utf-8")
                           if info.filename == output_path else output_zip.read(info.filename))
                result_zip.writestr(info, content)
        return result.getvalue()
