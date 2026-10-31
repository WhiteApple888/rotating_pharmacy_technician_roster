"""End-to-end checks for repairing the uploaded roster without a pattern."""

from datetime import date
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch

from ortools.sat.python import cp_model

from openpyxl import load_workbook

from roster_audit.individual_staff_constraints import DPW_CODES
from roster_audit.optimizer import (
    NoFeasibleRosterError, _change_cost, _count_changes, _marker,
    optimize_full_roster,
)
from roster_audit.repair_workbook import PRIORITY_COLUMNS, _load_priorities, write_repaired_roster
from roster_audit.service import audit_loaded_roster
from roster_audit.workbook import load_last_roster_rest_days, load_roster


SOURCE = Path(__file__).resolve().parents[2] / "sample" / "13Apr-10May_26.xlsx"
SHEET = "13Aprto10May"
START = date(2026, 4, 13)


def workbook_with_changes(changes: dict[str, str]) -> bytes:
    workbook = load_workbook(SOURCE)
    for cell, value in changes.items():
        workbook[SHEET][cell] = value
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def full_roster_with_leave() -> bytes:
    return workbook_with_changes({"Z6": "VL"})  # Staff 1, Day 15.


def full_roster_with_rest_day_leave() -> bytes:
    return workbook_with_changes({"AC6": "VL"})  # Staff 1, Day 18.


class FullOptimizerTests(unittest.TestCase):
    def optimize(self, contents):
        return optimize_full_roster(
            contents, start_date=START, public_holiday_days={19},
            worksheet_name=SHEET, max_time_seconds=30, fixed_staff={16},
        )

    def assert_audited_export(self, result):
        exported = load_roster(BytesIO(result.workbook_bytes), SHEET)
        prior = load_last_roster_rest_days(BytesIO(result.workbook_bytes), START)
        self.assertEqual(exported, result.roster)
        self.assertEqual(audit_loaded_roster(exported, {19}, START, prior.dates), [])

    def test_fixed_staff_and_restrictions_conflict_is_infeasible(self):
        original = load_roster(BytesIO(SOURCE.read_bytes()), SHEET)
        code = original.staff[0].shifts[0]
        with self.assertRaisesRegex(NoFeasibleRosterError, rf"Staff 1, day 1 is fixed to restricted shift {code}"):
            optimize_full_roster(SOURCE.read_bytes(), start_date=START,
                                 public_holiday_days={19}, worksheet_name=SHEET,
                                 fixed_staff={1}, shift_restrictions={1: {code}},
                                 max_time_seconds=10)

    def test_staff_count_reaches_workbook_and_solver(self):
        workbook = load_workbook(SOURCE)
        workbook["lastRosterRD"]["B18"] = workbook["lastRosterRD"]["B17"].value
        for column in range(1, 45):
            workbook[SHEET].cell(22, column).value = workbook[SHEET].cell(21, column).value
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        for count in (15, 17):
            with self.subTest(count=count), patch("roster_audit.optimizer._solve_roster", side_effect=RuntimeError("reached")) as solve:
                with self.assertRaisesRegex(RuntimeError, "reached"):
                    optimize_full_roster(output.getvalue(), start_date=START,
                                         public_holiday_days={19}, worksheet_name=SHEET,
                                         num_staff=count)
                self.assertEqual(len(solve.call_args.args[0].staff), count)
                self.assertEqual(len(solve.call_args.args[1]), count)

    def test_change_log_records_marker_and_base_changes(self):
        from roster_audit.individual_staff_constraints import Roster, StaffSchedule

        contents = workbook_with_changes({"D6": "OFF", "E6": "RD+", "F6": "NO+",
                                          "G6": "OFF+", "H6": "VL"})
        workbook = load_workbook(BytesIO(contents))
        workbook.create_sheet("Change Log")["A1"] = "obsolete"
        source = BytesIO()
        workbook.save(source)
        workbook.close()
        original = load_roster(BytesIO(source.getvalue()), SHEET)
        first = original.staff[0]
        updated = StaffSchedule(first.excel_row, ("OFF+", "RD", "NO++", "NO++", "VLRD") + first.shifts[5:])
        final = Roster(SHEET, (updated,) + original.staff[1:])
        self.assertEqual(_count_changes(original, final), (1, 5, 4))
        output = write_repaired_roster(source.getvalue(), SHEET, original, final, START, 1, 5, 4)
        workbook = load_workbook(BytesIO(output), data_only=True)
        log = workbook["Change Log"]
        self.assertEqual([log.cell(i, 2).value for i in range(1, 4)], [1, 5, 4])
        rows = list(log.iter_rows(min_row=6, values_only=True))
        self.assertEqual(len(rows), 5)
        self.assertEqual([(r[4], r[5], r[6], r[7], r[8]) for r in rows], [
            ("OFF", "OFF+", False, True, False),
            ("RD+", "RD", False, True, False),
            ("NO+", "NO++", False, True, False),
            ("OFF+", "NO++", True, True, True),
            ("VL", "VLRD", True, False, False),
        ])
        self.assertEqual((rows[0][0], rows[0][1], rows[0][2].date(), rows[0][3]),
                         (6, "D", START, workbook[SHEET]["A6"].value))
        workbook.close()

    def test_export_is_audited_after_serialization(self):
        contents = SOURCE.read_bytes()
        roster = load_roster(BytesIO(contents), SHEET)
        with patch("roster_audit.optimizer._solve_roster", return_value=(roster, cp_model.OPTIMAL)), \
             patch("roster_audit.optimizer.audit_loaded_roster", side_effect=[[], [object()]]) as audit:
            with self.assertRaisesRegex(NoFeasibleRosterError, "exported workbook did not pass"):
                self.optimize(contents)
        self.assertEqual(audit.call_count, 2)
        self.assertEqual(audit.call_args_list[1].args[0], roster)
        self.assertIsNot(audit.call_args_list[1].args[0], roster)

    def test_sample_priority_columns_follow_the_same_staff_rows(self):
        priorities = _load_priorities(SOURCE.read_bytes(), SHEET)
        self.assertEqual([next(staff for staff, rank in week.items() if rank == 1)
                          for week in priorities], [8, 7, 6, 5])

    def test_coverage_gap_can_be_repaired_without_reconstructing_a_leave_duty(self):
        result = self.optimize(workbook_with_changes({"Z6": "Z1"}))
        self.assertEqual(sum(staff.shifts[14] in DPW_CODES for staff in result.roster.staff), 4)
        self.assert_audited_export(result)

    def test_two_night_leave_cells_remain_leave_with_daily_coverage_restored(self):
        result = self.optimize(workbook_with_changes({"AD10": "VL", "AD16": "VL"}))
        for staff in (5, 11):
            self.assertIn(result.roster.staff[staff - 1].shifts[18], {"VL", "VLOFF", "VLRD"})
        self.assertEqual(sum(staff.shifts[18] in DPW_CODES for staff in result.roster.staff), 3)
        self.assert_audited_export(result)

    def test_no_xlsx_is_returned_when_leave_exceeds_daily_limit(self):
        with self.assertRaises(NoFeasibleRosterError):
            self.optimize(workbook_with_changes({"D8": "VL"}))

    def test_rest_day_leave_and_free_conversions_have_zero_penalized_changes(self):
        contents = full_roster_with_rest_day_leave()
        result = self.optimize(contents)
        original = load_roster(BytesIO(contents), SHEET)
        self.assertEqual(result.penalized_change_count, 0)
        self.assertGreater(result.cell_edit_count, 0)
        self.assertEqual(result.roster.staff[-1], original.staff[-1])
        self.assert_audited_export(result)

    def test_priorities_are_optional_and_need_not_align_with_a_pattern(self):
        workbook = load_workbook(BytesIO(full_roster_with_rest_day_leave()))
        for row in range(6, 22):
            for column in PRIORITY_COLUMNS:
                workbook[SHEET].cell(row, column).value = None
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        result = self.optimize(output.getvalue())
        self.assertEqual(result.penalized_change_count, 0)
        self.assert_audited_export(result)

    def test_holiday_shortage_is_repaired_from_uploaded_coverage(self):
        result = self.optimize(workbook_with_changes({"AD11": "VL"}))
        self.assertEqual(sum(staff.shifts[18] in {"XE", "XE1"} for staff in result.roster.staff), 2)
        self.assert_audited_export(result)

    def test_export_counts_use_uploaded_roster_and_exclude_free_conversions(self):
        contents = full_roster_with_leave()
        original = load_roster(BytesIO(contents), SHEET)
        result = self.optimize(contents)
        pairs = [(before, after) for old, new in zip(original.staff, result.roster.staff)
                 for before, after in zip(old.shifts, new.shifts)]
        self.assertEqual(result.penalized_change_count, sum(_change_cost(old, new) for old, new in pairs))
        self.assertEqual(result.cell_edit_count, sum(old != new for old, new in pairs))
        self.assertEqual(result.plus_change_count, sum(_marker(old) != _marker(new) for old, new in pairs))
        self.assertGreater(result.penalized_change_count, 0)
        # Priority 1 is used for the new night duty despite alternative covers
        # potentially requiring fewer paid edits elsewhere in the roster.
        self.assertEqual([
            staff for staff, (old, new) in enumerate(zip(original.staff, result.roster.staff), 1)
            if old.shifts[14] not in DPW_CODES and new.shifts[14] in DPW_CODES
        ], [6])
        self.assertIn(result.roster.staff[0].shifts[14], {"VL", "VLOFF", "VLRD"})
        uploaded = load_workbook(BytesIO(contents), read_only=True)
        self.assertEqual(uploaded[SHEET]["Z6"].value, "VL")
        uploaded.close()
        self.assert_audited_export(result)


if __name__ == "__main__":
    unittest.main()


class PlannerChoiceRegressionTests(unittest.TestCase):
    def test_exact_restrictions_and_defaults(self):
        from roster_audit.planner_choices import PlannerChoices
        choices = PlannerChoices.create(shift_restrictions={1: ["OFF+"]})
        self.assertEqual(choices.num_staff, 16)
        self.assertEqual(choices.fixed_staff, frozenset())
        self.assertEqual(choices.allowed(1, ("OFF", "OFF+", "OFF++")), ("OFF", "OFF++"))
        self.assertEqual(choices.allowed(2, ("OFF", "OFF+")), ("OFF", "OFF+"))
        self.assertEqual(PlannerChoices.create(shift_restrictions={1: ["OFF"]}).allowed(
            1, ("OFF", "OFF+", "OFF++")), ("OFF+", "OFF++"))

    def test_fixed_formula_without_cached_result_is_rejected(self):
        contents = workbook_with_changes({"D6": '= "OFF"'})
        with self.assertRaisesRegex(ValueError, "D6.*no readable calculated shift value"):
            optimize_full_roster(contents, start_date=START, public_holiday_days={19},
                                 worksheet_name=SHEET, fixed_staff={1}, max_time_seconds=10)

    def test_export_preserves_fixed_formula_cache_and_formatting(self):
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        from copy import copy
        from roster_audit.repair_workbook import load_repair_input
        from roster_audit.parsing import day_columns

        workbook = load_workbook(BytesIO(SOURCE.read_bytes()))
        sheet = workbook[SHEET]
        cell = sheet.cell(6, day_columns()[0])
        original_code = cell.value
        cell.value = '= "' + original_code + '"'
        cell.font = copy(sheet["E6"].font)
        expected_style = cell.style_id
        raw = BytesIO()
        workbook.save(raw)
        workbook.close()
        source = BytesIO()
        with ZipFile(BytesIO(raw.getvalue())) as src, ZipFile(source, "w") as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    formula = root.find('.//m:c[@r="D6"]', ns)
                    formula.set("t", "str")
                    formula.find("m:v", ns).text = original_code
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                dst.writestr(info, data)
        contents = source.getvalue()
        roster, _, _ = load_repair_input(contents, START, SHEET, 16, frozenset({1}))
        output = write_repaired_roster(contents, SHEET, roster, roster, START, 0, 0, 0, frozenset({1}))
        saved = load_workbook(BytesIO(output), data_only=False)
        self.assertEqual(saved[SHEET]["D6"].value, '= "' + original_code + '"')
        self.assertEqual(saved[SHEET]["D6"].style_id, expected_style)
        saved.close()
        saved = load_workbook(BytesIO(output), data_only=True)
        self.assertEqual(saved[SHEET]["D6"].value, original_code)
        saved.close()
