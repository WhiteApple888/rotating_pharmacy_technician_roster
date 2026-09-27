"""Tests for audit workbook report creation."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from roster_audit.individual_staff_constraints import AuditFailure, Roster
from roster_audit.report import write_report


class ReportTests(unittest.TestCase):
    def test_write_report_creates_summary_and_failure_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            report_path = Path(directory) / "audit.xlsx"
            write_report(
                report_path,
                [AuditFailure("Daily coverage", "coverage failed", day=3)],
                Path("source.xlsx"),
                Roster("April roster", ()),
                {3, 1},
            )

            workbook = load_workbook(report_path, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Summary", "Failures"])
                summary = workbook["Summary"]
                failures = workbook["Failures"]
                self.assertEqual(summary["A2"].value, "FAIL")
                self.assertEqual(summary["D2"].value, "1, 3")
                self.assertEqual(summary["E2"].value, 1)
                self.assertEqual(failures["A2"].value, "Daily coverage")
                self.assertEqual(failures["C2"].value, 3)
            finally:
                workbook.close()
