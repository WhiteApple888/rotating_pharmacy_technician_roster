"""Tests for Excel roster input."""

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from roster_audit.workbook import load_last_roster_rest_days, load_roster


class WorkbookTests(unittest.TestCase):
    def test_load_roster_normalises_the_fixed_grid(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "roster.xlsx"
            workbook = Workbook()
            workbook.active.cell(row=6, column=4, value=" xe1 ")
            workbook.save(path)
            workbook.close()

            roster = load_roster(path, None)

        self.assertEqual(len(roster.staff), 16)
        self.assertEqual(roster.staff[0].shifts[0], "XE1")

    def test_load_last_roster_rest_days_preserves_each_staff_members_prior_date(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "roster.xlsx"
            workbook = Workbook()
            history = workbook.create_sheet("lastRosterRD")
            for row in range(2, 18):
                history.cell(row=row, column=2, value=date(2026, 4, 12))
            history["B3"] = date(2026, 4, 10)
            workbook.save(path)
            workbook.close()

            rest_days = load_last_roster_rest_days(path, date(2026, 4, 13))

        self.assertEqual(len(rest_days.dates), 16)
        self.assertEqual(rest_days.dates[0], date(2026, 4, 12))
        self.assertEqual(rest_days.dates[1], date(2026, 4, 10))

    def test_load_last_roster_rest_days_rejects_a_date_on_or_after_day_one(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "roster.xlsx"
            workbook = Workbook()
            history = workbook.create_sheet("lastRosterRD")
            for row in range(2, 18):
                history.cell(row=row, column=2, value=date(2026, 4, 12))
            history["B2"] = date(2026, 4, 13)
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, r"lastRosterRD!B2 must be before roster Day 1"):
                load_last_roster_rest_days(path, date(2026, 4, 13))

    def test_load_last_roster_rest_days_rejects_non_dates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "roster.xlsx"
            workbook = Workbook()
            history = workbook.create_sheet("lastRosterRD")
            for row in range(2, 18):
                history.cell(row=row, column=2, value=date(2026, 4, 12))
            history["B2"] = "not a date"
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, r"lastRosterRD!B2 must contain an Excel date"):
                load_last_roster_rest_days(path, date(2026, 4, 13))
